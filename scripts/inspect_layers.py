#!/usr/bin/env python3
"""Inspect PPTX layers, extract full embedded pictures, and create a visual QA copy.

Requires python-pptx and Pillow. Does not perform OCR or slide rendering.
"""

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def walk(shapes, prefix=""):
    for shape in shapes:
        path = f"{prefix}/{shape.name}" if prefix else shape.name
        yield shape, path, bool(prefix)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes, path)


def safe_name(value):
    return re.sub(r"[^\w.-]+", "_", value, flags=re.UNICODE).strip("._")[:70] or "image"


def native_text(shape):
    if shape.has_text_frame:
        return shape.text
    if shape.has_table:
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
    return ""


def prune_for_isolation(shapes, excluded):
    """Keep picture objects, including pictures inside groups; never edit pixels."""
    for shape in list(shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            prune_for_isolation(shape.shapes, excluded)
            keep = len(shape.shapes) > 0
        else:
            keep = shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name not in excluded
        if not keep:
            element = shape._element
            element.getparent().remove(element)


def inspect(args):
    source = args.pptx.resolve()
    if args.isolation_pptx and args.isolation_pptx.resolve() == source:
        raise ValueError("Isolation output must differ from the source PPTX")
    if args.report and args.report.resolve() == source:
        raise ValueError("JSON report must differ from the source PPTX")
    with ZipFile(source) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"Corrupt ZIP entry: {bad}")
    prs = Presentation(source)
    if args.extract_dir:
        args.extract_dir.mkdir(parents=True, exist_ok=True)
    result = {
        "source": str(source),
        "zip_integrity": "ok",
        "slide_count": len(prs.slides),
        "canvas_inches": [round(prs.slide_width / 914400, 4), round(prs.slide_height / 914400, 4)],
        "visual_review_required": [
            "Render and compare the full slide with its reference",
            "Inspect each extracted full image, beyond its PowerPoint crop",
            "Check transparent assets against both light and dark backgrounds",
            "Render the isolation PPTX and inspect for residual text or layout fragments",
        ],
        "limitations": "No OCR or visual acceptance. PPI/coverage inside groups uses local coordinates.",
        "slides": [],
    }
    for number, slide in enumerate(prs.slides, 1):
        types = Counter()
        page = {"page": number, "text_objects": [], "pictures": [], "warnings": []}
        theme_effects = []
        for shape, path, grouped in walk(slide.shapes):
            types[str(shape.shape_type)] += 1
            value = native_text(shape)
            if value.strip():
                page["text_objects"].append({"name": path, "text": value})
            style = shape._element.find(qn("p:style"))
            effect = style.find(qn("a:effectRef")) if style is not None else None
            if effect is not None and effect.get("idx") not in (None, "0"):
                theme_effects.append({"name": path, "index": effect.get("idx")})
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            data = shape.image
            width, height = data.size
            fractions = [shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom]
            visible_w = max(0, 1 - fractions[0] - fractions[2])
            visible_h = max(0, 1 - fractions[1] - fractions[3])
            ppi = [
                round(width * visible_w / (shape.width / 914400), 1) if shape.width else None,
                round(height * visible_h / (shape.height / 914400), 1) if shape.height else None,
            ]
            coverage = shape.width * shape.height / (prs.slide_width * prs.slide_height)
            item = {
                "name": path,
                "sha256": hashlib.sha256(data.blob).hexdigest(),
                "original_pixels": [width, height],
                "crop_fractions_ltrb": [round(x, 5) for x in fractions],
                "effective_ppi_xy": ppi,
                "frame_area_fraction": round(coverage, 4),
                "inside_group": grouped,
            }
            if args.extract_dir:
                index = len(page["pictures"]) + 1
                dest = args.extract_dir / f"p{number:03d}_{index:03d}_{safe_name(shape.name)}.{data.ext}"
                dest.write_bytes(data.blob)
                item["extracted_file"] = str(dest.resolve())
                try:
                    with Image.open(dest) as image:
                        item["mode"] = image.mode
                        if "A" in image.getbands():
                            item["alpha_extrema"] = list(image.getchannel("A").getextrema())
                except OSError:
                    item["image_inspection_note"] = "Format requires a separate visual renderer"
            if coverage >= 0.8 and not grouped:
                page["warnings"].append(f"Review large raster frame: {path}")
            if any(x and x < 120 for x in ppi):
                page["warnings"].append(f"Review low effective resolution: {path} ({ppi} PPI)")
            if any(fractions):
                item["review_note"] = "Full image inspection required; crop hides part of the embedded file"
            page["pictures"].append(item)
        page["shape_types"] = dict(types)
        page["theme_effect_references"] = theme_effects
        if page["pictures"] and not page["text_objects"]:
            page["warnings"].append("No native text detected; check if this is intentionally an image-only page")
        result["slides"].append(page)
    if args.isolation_pptx:
        missing = set(args.exclude_picture_name)
        present = {s.name for slide in prs.slides for s, _, _ in walk(slide.shapes)
                   if s.shape_type == MSO_SHAPE_TYPE.PICTURE}
        result["unknown_excluded_names"] = sorted(missing - present)
        for slide in prs.slides:
            prune_for_isolation(slide.shapes, missing)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(args.background)
        args.isolation_pptx.parent.mkdir(parents=True, exist_ok=True)
        prs.save(args.isolation_pptx)
        result["isolation_pptx"] = str(args.isolation_pptx.resolve())
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(payload + "\n", encoding="utf-8")
        print(json.dumps({"report": str(args.report.resolve()),
                          "slides": result["slide_count"],
                          "pictures": sum(len(p["pictures"]) for p in result["slides"]),
                          "native_text_objects": sum(len(p["text_objects"]) for p in result["slides"]),
                          "visual_review": "required"}, ensure_ascii=False))
    else:
        print(payload)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--extract-dir", type=Path)
    parser.add_argument("--isolation-pptx", type=Path)
    parser.add_argument("--exclude-picture-name", action="append", default=[])
    parser.add_argument("--background", default="FFFFFF", help="Six-digit RGB hex for the QA slide background")
    args = parser.parse_args()
    args.background = args.background.lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", args.background):
        parser.error("--background must be a six-digit RGB hex value")
    try:
        inspect(args)
    except (OSError, ValueError) as error:
        parser.exit(1, f"Inspection failed: {error}\n")


if __name__ == "__main__":
    main()
